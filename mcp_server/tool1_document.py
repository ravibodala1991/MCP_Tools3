"""
Tool 1 – Document Upload, Parsing, Embeddings & Summarization
=============================================================
Flow:
  1. Receive a base64-encoded file (PDF / TXT / DOCX) + metadata.
  2. Upload the raw file to S3 (raw-docs bucket).
  3. Parse text with PyMuPDF (PDF) or python-docx (DOCX) or plain read (TXT).
  4. Chunk the text and generate embeddings via Amazon Titan Text Embeddings v2.
  5. Store each chunk + embedding in OpenSearch Serverless (vector index).
  6. Summarize the full document with Amazon Bedrock Claude 3 Sonnet.
  7. Return a structured result with doc_id, page_count, chunk_count, summary.
"""

from __future__ import annotations

import base64
import hashlib
import io
import json
import os
import time
import uuid
from typing import Any

import boto3
from langchain_aws import BedrockEmbeddings
from langchain_aws.chat_models import ChatBedrock
from langchain_community.document_loaders import PyMuPDFLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from opensearchpy import OpenSearch, RequestsHttpConnection
from requests_aws4auth import AWS4Auth

# ── env / config ──────────────────────────────────────────────────────────────
AWS_REGION = os.environ.get("AWS_REGION", "us-east-1")
S3_RAW_BUCKET = os.environ.get("S3_RAW_BUCKET", "mcp-raw-docs")
OPENSEARCH_ENDPOINT = os.environ.get("OPENSEARCH_ENDPOINT", "")  # host only, no https://
OPENSEARCH_INDEX = os.environ.get("OPENSEARCH_INDEX", "mcp-documents")
EMBED_MODEL_ID = os.environ.get("EMBED_MODEL_ID", "amazon.titan-embed-text-v2:0")
LLM_MODEL_ID = os.environ.get("LLM_MODEL_ID", "anthropic.claude-3-sonnet-20240229-v1:0")
CHUNK_SIZE = int(os.environ.get("CHUNK_SIZE", "1000"))
CHUNK_OVERLAP = int(os.environ.get("CHUNK_OVERLAP", "200"))

# ── AWS clients ───────────────────────────────────────────────────────────────
s3_client = boto3.client("s3", region_name=AWS_REGION)
bedrock_runtime = boto3.client("bedrock-runtime", region_name=AWS_REGION)
credentials = boto3.Session().get_credentials()
awsauth = AWS4Auth(
    credentials.access_key,
    credentials.secret_key,
    AWS_REGION,
    "aoss",
    session_token=credentials.token,
)


def _get_opensearch_client() -> OpenSearch:
    """Return an authenticated OpenSearch Serverless client."""
    return OpenSearch(
        hosts=[{"host": OPENSEARCH_ENDPOINT, "port": 443}],
        http_auth=awsauth,
        use_ssl=True,
        verify_certs=True,
        connection_class=RequestsHttpConnection,
        timeout=30,
    )


def _ensure_index(client: OpenSearch, dims: int = 1536) -> None:
    """Create the vector index if it does not exist."""
    if client.indices.exists(index=OPENSEARCH_INDEX):
        return
    body = {
        "settings": {"index": {"knn": True}},
        "mappings": {
            "properties": {
                "doc_id": {"type": "keyword"},
                "filename": {"type": "keyword"},
                "page": {"type": "integer"},
                "chunk_index": {"type": "integer"},
                "text": {"type": "text"},
                "embedding": {
                    "type": "knn_vector",
                    "dimension": dims,
                    "method": {
                        "name": "hnsw",
                        "space_type": "cosinesimil",
                        "engine": "nmslib",
                    },
                },
            }
        },
    }
    client.indices.create(index=OPENSEARCH_INDEX, body=body)


def _parse_file(file_bytes: bytes, filename: str) -> list[Document]:
    """Parse raw bytes into a list of LangChain Documents (one per page/chunk)."""
    ext = filename.rsplit(".", 1)[-1].lower()

    if ext == "pdf":
        # Write to a temp file so PyMuPDFLoader can open it
        tmp_path = f"/tmp/{uuid.uuid4()}.pdf"
        with open(tmp_path, "wb") as f:
            f.write(file_bytes)
        loader = PyMuPDFLoader(tmp_path)
        docs = loader.load()
        os.remove(tmp_path)
        return docs

    if ext in ("txt", "md"):
        text = file_bytes.decode("utf-8", errors="replace")
        return [Document(page_content=text, metadata={"page": 0, "source": filename})]

    if ext == "docx":
        try:
            import docx  # python-docx

            doc = docx.Document(io.BytesIO(file_bytes))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            return [Document(page_content=text, metadata={"page": 0, "source": filename})]
        except ImportError:
            raise RuntimeError("python-docx is required for .docx files. pip install python-docx")

    raise ValueError(f"Unsupported file type: .{ext}. Supported: pdf, txt, md, docx")


def _embed_texts(texts: list[str]) -> list[list[float]]:
    """Batch-embed texts using Amazon Titan Embeddings v2."""
    embedder = BedrockEmbeddings(
        client=bedrock_runtime,
        model_id=EMBED_MODEL_ID,
        region_name=AWS_REGION,
    )
    return embedder.embed_documents(texts)


def _summarize(full_text: str) -> str:
    """Summarize the document using Claude 3 Sonnet via Bedrock."""
    llm = ChatBedrock(
        client=bedrock_runtime,
        model_id=LLM_MODEL_ID,
        region_name=AWS_REGION,
        model_kwargs={"max_tokens": 1024, "temperature": 0.2},
    )
    prompt = (
        "You are a document analyst. Provide a concise, structured summary of the "
        "following document. Include: main topic, key points, and any action items.\n\n"
        f"DOCUMENT:\n{full_text[:12000]}"  # cap to avoid token limits
    )
    response = llm.invoke(prompt)
    return response.content


# ── Main tool function ────────────────────────────────────────────────────────

def upload_parse_embed_summarize(
    file_b64: str,
    filename: str,
    metadata: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """
    Upload, parse, embed, and summarize a document.

    Args:
        file_b64:  Base64-encoded file content.
        filename:  Original filename including extension (pdf / txt / docx / md).
        metadata:  Optional key-value metadata to attach to every chunk.

    Returns:
        dict with doc_id, filename, page_count, chunk_count, s3_key, summary.
    """
    metadata = metadata or {}
    doc_id = hashlib.sha256(f"{filename}{time.time()}".encode()).hexdigest()[:16]

    # 1. Decode
    file_bytes = base64.b64decode(file_b64)

    # 2. Upload raw file to S3
    s3_key = f"raw-docs/{doc_id}/{filename}"
    s3_client.put_object(
        Bucket=S3_RAW_BUCKET,
        Key=s3_key,
        Body=file_bytes,
        Metadata={"doc_id": doc_id, **{k: str(v) for k, v in metadata.items()}},
    )

    # 3. Parse
    raw_docs = _parse_file(file_bytes, filename)
    page_count = len(raw_docs)
    full_text = "\n\n".join(d.page_content for d in raw_docs)

    # 4. Chunk
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
    )
    chunks = splitter.split_documents(raw_docs)

    # 5. Embed
    texts = [c.page_content for c in chunks]
    embeddings = _embed_texts(texts)

    # 6. Index into OpenSearch
    os_client = _get_opensearch_client()
    _ensure_index(os_client, dims=len(embeddings[0]) if embeddings else 1536)

    bulk_body: list[dict] = []
    for i, (chunk, emb) in enumerate(zip(chunks, embeddings)):
        bulk_body.append({"index": {"_index": OPENSEARCH_INDEX, "_id": f"{doc_id}_{i}"}})
        bulk_body.append(
            {
                "doc_id": doc_id,
                "filename": filename,
                "page": chunk.metadata.get("page", 0),
                "chunk_index": i,
                "text": chunk.page_content,
                "embedding": emb,
                **metadata,
            }
        )
    if bulk_body:
        os_client.bulk(body=bulk_body, refresh=True)

    # 7. Summarize
    summary = _summarize(full_text)

    return {
        "doc_id": doc_id,
        "filename": filename,
        "s3_key": s3_key,
        "page_count": page_count,
        "chunk_count": len(chunks),
        "summary": summary,
    }
